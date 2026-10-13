"""
Shared design system for The Media Buyer's Client Reporting Kit.
Used by build_kpi_calculator.py, build_dashboard.py, build_client_report.py

Edit the palette / CURRENCY here once and regenerate every workbook.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------- brand config
KIT_NAME      = "The Media Buyer's Client Reporting Kit"
CURRENCY      = "EGP"          # <- change once, then run all three build scripts
CURRENCY_CODE = "EGP"

# ---------------------------------------------------------------------- palette
INK        = "111827"
INK_SOFT   = "374151"
MUTED      = "6B7280"
HAIRLINE   = "E5E7EB"
PANEL      = "F3F4F6"
PANEL_DK   = "E5E7EB"
WHITE      = "FFFFFF"
ACCENT     = "F97316"   # signal orange
ACCENT_DK  = "C2410C"
ACCENT_BG  = "FFF7ED"
INPUT_BG   = "FFF7E6"   # amber = "type here"
CALC_BG    = "FFFFFF"   # white  = "calculated, leave alone"
GOOD       = "15803D";  GOOD_BG = "DCFCE7"
WARN       = "B45309";  WARN_BG = "FEF3C7"
BAD        = "B91C1C";  BAD_BG  = "FEE2E2"
INFO       = "1D4ED8";  INFO_BG = "DBEAFE"

# --------------------------------------------------------------- number formats
F_INT   = '#,##0'
F_2D    = '#,##0.00'
F_PCT   = '0.00%'
F_PCT1  = '0.0%'
F_RATIO = '0.00"x"'
F_RATIO4= '0.000"x"'
F_DATE  = 'yyyy-mm-dd'

def money(sym=None):
    sym = sym or CURRENCY_CODE
    return f'#,##0.00" {sym}"'

def money0(sym=None):
    sym = sym or CURRENCY_CODE
    return f'#,##0" {sym}"'

# --------------------------------------------------------------------- fonts
def f_title(sz=16):   return Font(name="Inter", size=sz, bold=True, color=INK)
def f_h1(sz=12):      return Font(name="Inter", size=sz, bold=True, color=INK)
def f_h2(sz=10.5):    return Font(name="Inter", size=sz, bold=True, color=INK)
def f_body(sz=10):    return Font(name="Inter", size=sz, color=INK_SOFT)
def f_small(sz=9):    return Font(name="Inter", size=sz, color=MUTED)
def f_tiny(sz=8):     return Font(name="Inter", size=sz, color=MUTED)
def f_italic(sz=9.5): return Font(name="Inter", size=sz, italic=True, color=MUTED)
def f_section(sz=10.5): return Font(name="Inter", size=sz, bold=True, color=WHITE)
def f_input(sz=10):   return Font(name="Inter", size=sz, bold=True, color=ACCENT_DK)
def f_calc(sz=10):    return Font(name="Inter", size=sz, bold=True, color=INK)
def f_kpi(sz=18):     return Font(name="Inter", size=sz, bold=True, color=INK)
def f_kpi_lbl(sz=9):  return Font(name="Inter", size=sz, bold=False, color=MUTED)

# --------------------------------------------------------------------- fills
def fill(hexcolor):   return PatternFill("solid", fgColor=hexcolor)
F_SECTION  = fill(INK)
F_HEAD     = fill(PANEL_DK)
F_PANEL    = fill(PANEL)
F_INPUT    = fill(INPUT_BG)
F_CALC     = fill(CALC_BG)
F_ACCENT   = fill(ACCENT)
F_ACCENTBG = fill(ACCENT_BG)
F_GOODBG   = fill(GOOD_BG)
F_WARNBG   = fill(WARN_BG)
F_BADBG    = fill(BAD_BG)
F_INFOBG   = fill(INFO_BG)
F_WHITE    = fill(WHITE)

# --------------------------------------------------------------------- borders
_s = Side(style="thin", color=HAIRLINE)
_sa = Side(style="thin", color=ACCENT)
def b_all(color=HAIRLINE):
    sd = Side(style="thin", color=color)
    return Border(left=sd, right=sd, top=sd, bottom=sd)
def b_box(color=INK):
    sd = Side(style="thin", color=color)
    return Border(left=sd, right=sd, top=sd, bottom=sd)
B_BOX   = Border(left=Side(style="thin", color=INK), right=Side(style="thin", color=INK),
                 top=Side(style="thin", color=INK), bottom=Side(style="thin", color=INK))
B_INPUT = Border(left=_sa, right=_sa, top=_sa, bottom=_sa)
B_BOTTOM= Border(bottom=Side(style="thin", color=HAIRLINE))

# ------------------------------------------------------------------- alignment
A_L  = Alignment(horizontal="left",   vertical="center", wrap_text=False)
A_LW = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
A_R  = Alignment(horizontal="right",  vertical="center")
A_C  = Alignment(horizontal="center", vertical="center")
A_CW = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LI = Alignment(horizontal="left",   vertical="center", indent=1)


# ------------------------------------------------------------------- utilities
def put(ws, ref, value, font=None, fill_=None, fmt=None, align=None, border=None,
        row_h=None, col_w=None, literal=False):
    """Write one cell with optional styling. Returns the cell.

    literal=True keeps a string that starts with "=" as TEXT. openpyxl would
    otherwise store it as a real formula (<f>), and Google Sheets evaluates
    those on import -- so documentation text becomes a #REF!/#ERROR! cell.
    """
    c = ws[ref]
    c.value = value
    if literal and isinstance(value, str) and value.lstrip().startswith("="):
        c.data_type = "s"
    if font:   c.font = font
    if fill_:  c.fill = fill_
    if fmt:    c.number_format = fmt
    if align:  c.alignment = align
    if border: c.border = border
    if row_h is not None:
        ws.row_dimensions[c.row].height = row_h
    if col_w is not None:
        ws.column_dimensions[c.column_letter].width = col_w
    return c


def autofit_row(ws, ref, text, font_pt=8.5, pad=6, min_h=20):
    """Set a row tall enough that the text in `ref` is not clipped.

    Excel column widths are in units of the default 11pt font; Consolas at
    8.5pt fits roughly 1.3x that many characters per line. Wraps each explicit
    line to estimate the displayed line count, then converts to points.
    """
    import textwrap
    col = "".join(ch for ch in ref if ch.isalpha())
    width = ws.column_dimensions[col].width or 9
    per_line = max(10, int(width * 11.0 / font_pt * 0.95))
    lines = sum(max(1, len(textwrap.wrap(ln, per_line)) or 1)
                for ln in str(text).split("\n"))
    ws.row_dimensions[int("".join(ch for ch in ref if ch.isdigit()))].height = max(
        min_h, lines * (font_pt * 1.45) + pad)


def put_row(ws, row, col_start, values, font=None, fill_=None, fmt=None,
            align=None, border=None):
    out = []
    for i, v in enumerate(values):
        ref = f"{chr(ord('A') + col_start - 1 + i)}{row}" if col_start <= 26 else None
        if ref is None:
            from openpyxl.utils import get_column_letter
            ref = f"{get_column_letter(col_start + i)}{row}"
        out.append(put(ws, ref, v, font=font, fill_=fill_, fmt=fmt,
                       align=align, border=border))
    return out


def section_bar(ws, row, first_col, last_col, text):
    """Dark full-width section header bar."""
    from openpyxl.utils import get_column_letter
    a, b = get_column_letter(first_col), get_column_letter(last_col)
    ws.merge_cells(f"{a}{row}:{b}{row}")
    c = ws[f"{a}{row}"]
    c.value = text
    c.font = f_section()
    c.fill = F_SECTION
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for i in range(first_col, last_col + 1):
        ws.cell(row=row, column=i).fill = F_SECTION
    ws.row_dimensions[row].height = 22
    return c


def title_block(ws, title, subtitle=None, currency=None, last_col=7):
    from openpyxl.utils import get_column_letter
    lc = get_column_letter(last_col)
    ws.merge_cells(f"A1:{lc}1")
    c = ws["A1"]; c.value = title; c.font = f_title(16); c.alignment = A_L
    ws.row_dimensions[1].height = 26
    if subtitle:
        ws.merge_cells(f"A2:{lc}2")
        c = ws["A2"]; c.value = subtitle; c.font = f_italic(9.5)
        c.alignment = A_L
        ws.row_dimensions[2].height = 15
    if currency:
        ws.merge_cells(f"A3:{lc}3")
        c = ws["A3"]
        c.value = f"All amounts in {currency}. Change the currency label on 'Start Here' after you copy this file."
        c.font = f_small(9); c.alignment = A_L
        ws.row_dimensions[3].height = 14


def label_value_hdr(ws, row, col, label="Metric", value="Value", note="Notes"):
    """Small grey header row above a label/value block."""
    put(ws, f"{chr(64+col)}{row}",   label, f_h2(9.5), F_HEAD, align=A_LI)
    put(ws, f"{chr(65+col)}{row}",   value, f_h2(9.5), F_HEAD, align=A_C)
    put(ws, f"{chr(66+col)}{row}",   note,  f_h2(9.5), F_HEAD, align=A_LI)
    ws.row_dimensions[row].height = 17


def metric_row(ws, row, col, label, value, note="", fmt=F_2D, is_input=False,
               strong=False, border=True, note_font=None):
    """One label | value | note row."""
    from openpyxl.utils import get_column_letter
    c1, c2, c3 = (get_column_letter(col + i) for i in range(3))
    bd = b_all() if border else None
    put(ws, f"{c1}{row}", label, f_h2(10) if strong else f_body(10),
        F_INPUT if is_input else None, align=A_LI, border=bd)
    put(ws, f"{c2}{row}", value,
        f_input(10) if is_input else (f_calc(10) if strong else f_calc(10)),
        F_INPUT if is_input else None, fmt=fmt, align=A_R, border=bd)
    put(ws, f"{c3}{row}", note, note_font or f_small(8.5), align=A_LI, border=bd)
    return ws[f"{c2}{row}"]


def verdict_cell(ws, ref, text, font=None, fill_=None, align=None):
    c = ws[ref]
    c.value = text
    c.font = font or Font(name="Inter", size=10, bold=True, color=INK)
    c.fill = fill_ or F_CALC
    c.alignment = align or Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    return c


STATUS_FONTS = {
    "good": Font(name="Inter", size=10, bold=True, color=GOOD),
    "warn": Font(name="Inter", size=10, bold=True, color=WARN),
    "bad":  Font(name="Inter", size=10, bold=True, color=BAD),
    "info": Font(name="Inter", size=10, bold=True, color=INFO),
    "muted":Font(name="Inter", size=10, color=MUTED),
}
STATUS_FILLS = {"good": F_GOODBG, "warn": F_WARNBG, "bad": F_BADBG, "info": F_INFOBG,
                "muted": F_PANEL}


def std_widths(ws, widths):
    """widths: dict like {'A':38,'B':16,...}"""
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def gridlines_off(ws):
    ws.sheet_view.showGridLines = False


def kpi_card(ws, top_row, left_col, label, formula, fmt, sub=None, span=2):
    """
    A tidy KPI tile: label on top, big number, optional sub-line.
    Occupies `span` columns and 3 rows.
    """
    from openpyxl.utils import get_column_letter
    lc = get_column_letter(left_col)
    rc = get_column_letter(left_col + span - 1)
    r1, r2, r3 = top_row, top_row + 1, top_row + 2
    ws.merge_cells(f"{lc}{r1}:{rc}{r1}")
    c = ws[f"{lc}{r1}"]; c.value = label; c.font = f_kpi_lbl(9)
    c.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    c.fill = F_PANEL
    ws.merge_cells(f"{lc}{r2}:{rc}{r2}")
    c = ws[f"{lc}{r2}"]; c.value = formula; c.font = f_kpi(17)
    c.number_format = fmt
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    c.fill = F_PANEL
    ws.merge_cells(f"{lc}{r3}:{rc}{r3}")
    c = ws[f"{lc}{r3}"]; c.value = sub or ""
    c.font = f_tiny(8)
    c.alignment = Alignment(horizontal="left", vertical="top", indent=1)
    c.fill = F_PANEL
    for r in (r1, r2, r3):
        for i in range(left_col, left_col + span):
            ws.cell(row=r, column=i).fill = F_PANEL
    ws.row_dimensions[r1].height = 14
    ws.row_dimensions[r2].height = 24
    ws.row_dimensions[r3].height = 12
    return ws[f"{lc}{r2}"]


def rep_bar(maxwidth=18):
    """In-cell bar formula fragment: wrap your 0-1 ratio in this."""
    return f'=REPT("█", ROUND({maxwidth}*{{ratio}},0))'
